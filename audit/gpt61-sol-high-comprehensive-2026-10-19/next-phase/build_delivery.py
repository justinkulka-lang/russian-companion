from pathlib import Path
import csv,io,json,hashlib,zipfile,base64,textwrap,shutil
from docx import Document
r=Path(__file__).resolve().parents[3];base=r/'audit/gpt61-sol-high-comprehensive-2026-10-19';n=base/'next-phase';art=Path('/workspace/artifacts');results=base/'continuation-results/2026-10-19-199b3229'
prompt='''Continue in this SAME ChatGPT GPT-6.1 Sol / High audit conversation. Preserve your completed work and the thirteen findings CC001–CC013; new findings start at CC014.

Repository: justinkulka-lang/russian-companion
Branch: cursor/content-audit-corrections-964f
Existing Draft PR64
Audited commit: 5c176eb22f4a93bf5f7fd20af15e79ec8c4748ed
Actual target tree: f95bfa3c1ae6b8ac07b42b6db9fa69c5e247830f
Verified engineering receipt: 5322a3dc83eaec5dabfdad555afffcb820d11a40

Read the latest results under:
audit/gpt61-sol-high-comprehensive-2026-10-19/continuation-results/2026-10-19-199b3229/

Read INDEPENDENT_AUDIT_REPORT.md, CODEX_RETURN_HANDOFF.md, RUNTIME_VERIFICATION_HANDOFF.md, NEXT_ENGINEERING_CHECKPOINT.md and CODEX_RECEIPT_VERIFICATION.json. All sixteen outputs and exact finding evidence are inside GPT61_SAME_CHAT_AUDIT_CONTINUATION_RESULTS.zip, including ledgers without standalone copies. Its size is 601,558 bytes and SHA256 is 199b3229665b769b043434f1d6a69abb242d91b16fd6ed26187fad1acb294185.

Read next-phase/DELIVERABLES.md, INPUTS_AND_BOUNDARIES.md and FINDING_WORK_QUEUE.csv in the same parent directory. The current verdict is INCOMPLETE: thirteen findings, five HIGH and eight MEDIUM. Original findings, failures and audit results must remain unchanged.

NEXT WORK

1. Prepare a concrete implementation work order for all thirteen findings. Specify exact entities and paths, authority/source constraints, the smallest justified correction, positive and negative tests, historical/digest/schema/assignment effects, prerequisites and remaining qualifications. Prioritize CC001/CC010, CC012/CC011 and CC013. CC008 needs actual Kotlin validator reproduction; CC009 needs qualified source review. Do not invent Russian, stress, audio or rights evidence, or copy private-book examples into proposed fixes.

2. Continue substantive unread independent inspection from your existing ledgers: applicable app code, complete teaching/check/remediation relations, full source contexts and rights, provenance, historical evidence and remaining book coverage. Do not restart completed work. Retain 6,427 target paths, 622 item IDs, 75 historical targets, eighteen IA identities and 12,034 entity rows. Update inspection status only after actual review. Hashing, OCR processing or generating a file is not semantic inspection.

3. Keep Android tests, compilation, lint, accessibility, performance and runtime behavior NOT_EXECUTED unless actually executed with the necessary tooling. A GitHub-only ChatGPT session cannot supply those results by reading source or logs. Produce the concrete engineering assignment in ENGINEERING_EXECUTOR_PROMPT.md for a separate Codex executor, with exact scenarios and evidence requests. This does not authorize source changes in this audit chat.

4. Return updated versions of all sixteen audit outputs, plus IMPLEMENTATION_WORK_ORDER.csv, EXECUTION_EVIDENCE_REQUESTS.csv and NEXT_ACTIONS.md. Follow next-phase/DELIVERABLES.md. Include checkpoint.json, SHA256SUMS.txt and an updated change summary. Publish audit artifacts only to PR64 if safe writes are supported, preserving concurrent work and never force-pushing. Otherwise return downloadable files and paste the work order and report in chat. Do not leave the only results in temporary storage.

The attached complete TXT packet, if supplied, contains the current audit, exact finding evidence and the full curriculum supplement. Use Python to decode it as directed. Reuse the current private Penguin PDF already attached here: 13,596,737 bytes; SHA256 2bcc0c45fac4508adf283f4dd939a7e578cf2c1306fdadf4966300532beb6303. If unavailable, request it again while continuing independent work. No PDF, extracted prose, exercises, scans or book-native examples in returned artifacts.

Use one agent. Preserve all 64 authentic histories, original failures, legacy policy/schema meanings and the unconfirmed Unit7 restore cause. Keep four screenshot exporters/16 methods excluded. Do not sum the supplied 1,225-case failed broad run and separate 26/31-method followups into a new all-pass suite.

A complete diagnostic may contain open defects, but unread required scope keeps the audit INCOMPLETE. No app fixes, merge, freeze, APK assembly/signing, formal approval-stage advancement or release approval in this audit chat. Physical Galaxy S26 Ultra airplane-mode/reboot acceptance remains LAST and OUTSTANDING after final comprehensive independent review of corrected source.
'''
(n/'SAME_CHAT_PROMPT.md').write_text(prompt)
spec='''# Next-phase deliverables

The same ChatGPT audit must preserve and update all sixteen outputs:

1. INDEPENDENT_AUDIT_REPORT.md
2. FINDINGS.csv
3. ITEM_ASSESSMENT_AUDIT.csv
4. UNIT_COVERAGE.csv
5. TARGET_75_DISPOSITIONS.csv
6. REQUIREMENTS_MATRIX.csv
7. SOURCE_AND_BINDING_REVIEW.json
8. AUDIT_SCOPE_AND_EVIDENCE.json
9. CODEX_RETURN_HANDOFF.md
10. APP_ENGINEERING_AUDIT.csv
11. REPOSITORY_FILE_COVERAGE.csv
12. CURRICULUM_ENTITY_COVERAGE.csv
13. FINDING_RECONCILIATION.csv
14. GAP_CLOSURE.csv
15. RUNTIME_VERIFICATION_HANDOFF.md
16. BOOK_COVERAGE_CONTINUATION.csv

Preserve original schemas and CC001–CC013, next CC014. Keep every existing row identity, including 12,034 entity rows. No invented passing or inspected status.

Additional outputs:

17. IMPLEMENTATION_WORK_ORDER.csv: finding_id,severity,current_disposition,exact_entities_and_paths,authority_and_source_constraints,bounded_correction,positive_control,negative_regression,history_and_identity_effects,prerequisites_or_blockers,executor,independent_recheck_needed. One row per thirteen findings, plus any new findings.
18. EXECUTION_EVIDENCE_REQUESTS.csv: requirement_id,scenario,required_environment,command_or_test_identity,expected_behavior,evidence_required,execution_status,remaining_blocker. Planned checks remain NOT_EXECUTED.
19. NEXT_ACTIONS.md: distinguish executable engineering, independent unread review, source/native/rights research and deferred physical acceptance. Give exact next actions and dependencies.

Include checkpoint.json, SHA256SUMS.txt and audit-change-summary.md. Preserve originals and return new outputs in a new archive, plus accessible text/Word when needed. Record actual model/effort verification basis, commit and tree separately, and supplied versus independently reproduced evidence.

The separate Codex engineering executor returns ENGINEERING_CHANGE_REPORT.md, FINDING_ENGINEERING_DISPOSITIONS.csv, VERIFICATION_MANIFEST.json and safe raw execution evidence, plus commits and the PR64 update. Record exact commands, input commits/hashes, selectors, outcomes including skipped/unrun methods, lint counts, gate attachment and source/rights blockers. Engineering results do not overwrite independent audit dispositions or establish formal acceptance.

No runtime or code correction was performed in preparing this handoff. Physical-device acceptance remains last and outstanding.
'''
(n/'DELIVERABLES.md').write_text(spec)
# Seed work queue is copied from audit evidence, not new engineering dispositions.
with zipfile.ZipFile(results/'GPT61_SAME_CHAT_AUDIT_CONTINUATION_RESULTS.zip') as z:
 findings=list(csv.DictReader(io.StringIO(z.read('FINDINGS.csv').decode('utf-8-sig'))))
with (n/'FINDING_WORK_QUEUE.csv').open('w',newline='') as f:
 keys=['finding_id','severity','category','unit','content_id','current_path','locator','requirement_id','observed_behavior','acceptance_test','audit_status','engineering_status']
 w=csv.DictWriter(f,fieldnames=keys,lineterminator='\n');w.writeheader()
 for item in findings:
  row={k:item.get(k,'') for k in keys};row['audit_status']=item['status'];row['engineering_status']='NOT_STARTED_NO_FIXES_IN_THIS_PACKET';w.writerow(row)
(n/'IMPLEMENTATION_WORK_ORDER_TEMPLATE.csv').write_text('finding_id,severity,current_disposition,exact_entities_and_paths,authority_and_source_constraints,bounded_correction,positive_control,negative_regression,history_and_identity_effects,prerequisites_or_blockers,executor,independent_recheck_needed\n')
(n/'EXECUTION_EVIDENCE_REQUESTS_TEMPLATE.csv').write_text('requirement_id,scenario,required_environment,command_or_test_identity,expected_behavior,evidence_required,execution_status,remaining_blocker\n')
word=Document();word.add_heading('Russian Companion: next audit and engineering handoff',0)
word.add_paragraph('Use Part 1 in the SAME ChatGPT audit conversation. Part 2 is for a separate Codex engineering executor with a real checkout and Android tooling. No code fixes or runtime verification were performed in preparing this document.')
word.add_heading('Part 1 — copy and paste into the same ChatGPT chat',1)
for line in prompt.splitlines():word.add_paragraph(line)
word.add_heading('Required deliverables',1)
for line in spec.splitlines():word.add_paragraph(line)
word.add_heading('Part 2 — separate Codex engineering assignment',1)
for line in (n/'ENGINEERING_EXECUTOR_PROMPT.md').read_text().splitlines():word.add_paragraph(line)
word.add_heading('Current finding register — unchanged audit evidence',1)
for item in findings:
 word.add_heading(item['finding_id']+' — '+item['severity']+' — '+item['category'],2)
 for label,key in [('Observed','observed_behavior'),('Paths','current_path'),('Acceptance test','acceptance_test'),('Audit status','status')]:word.add_paragraph(label+': '+item.get(key,''))
word_path=art/'Russian_Companion_Next_Steps_and_Prompts.docx';word.save(word_path)
# Portable complete payload includes the current immutable result archive and the lossless curriculum supplement.
entries={}
for p in n.iterdir():
 if p.is_file() and p.name != 'DELIVERY.json':entries['instructions/'+p.name]=p.read_bytes()
entries['current-audit/GPT61_SAME_CHAT_AUDIT_CONTINUATION_RESULTS.zip']=(results/'GPT61_SAME_CHAT_AUDIT_CONTINUATION_RESULTS.zip').read_bytes()
supp=r/'audit/gpt61-sol-high-comprehensive-2026-10-19/continuation/downloads/GPT61_Same_Chat_Audit_Continuation_5c176eb2.zip'
entries['curriculum-inputs/GPT61_Same_Chat_Audit_Continuation_5c176eb2.zip']=supp.read_bytes()
entries['README.txt']=b'Use instructions/SAME_CHAT_PROMPT.md in the same ChatGPT chat. Engineering prompt is for a separate Codex executor. Nested archives preserve original evidence. No PDF, SDK, Git database, APK or new runtime result is included.\n'
manifest={'auditTarget':'5c176eb22f4a93bf5f7fd20af15e79ec8c4748ed','receiptBaseline':'5322a3dc83eaec5dabfdad555afffcb820d11a40','newAuditPerformed':False,'sourceFixes':False,'files':[{'path':name,'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()} for name,data in sorted(entries.items())]}
entries['MANIFEST.json']=(json.dumps(manifest,indent=2)+'\n').encode()
archive=art/'Russian_Companion_Next_Phase_Packet.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=9) as z:
 for name,data in sorted(entries.items()):z.writestr(name,data)
with zipfile.ZipFile(archive) as z:
 assert z.testzip() is None
 for row in manifest['files']:assert hashlib.sha256(z.read(row['path'])).hexdigest()==row['sha256']
data=archive.read_bytes();digest=hashlib.sha256(data).hexdigest()
header=f'''COMPLETE RUSSIAN COMPANION NEXT-PHASE PACKET\n\nUse this in the SAME ChatGPT audit chat. The entire packet is embedded losslessly below. Use Python/file tools; do not attempt to read base64 as source evidence.\n\nRead this attachment as UTF-8. Split once on the exact marker line BEGIN_BASE64_NEXT_PHASE_PACKET. Base64-decode everything after that marker. Verify byte size {len(data)} and SHA256 {digest}. Open with zipfile.ZipFile(io.BytesIO(decoded)), check CRCs, and extract into a fresh directory after rejecting absolute or .. paths. Verify every MANIFEST.json entry hash/size. Open the nested current-audit and curriculum-inputs archives, preserving original files; follow instructions/SAME_CHAT_PROMPT.md. If Python tools are unavailable, use the private GitHub repository paths in the prompt and state any inaccessible inputs.\n\nPrivate Penguin PDF is separate; reuse the current verified attachment in this chat. This packet contains no PDF or new Android execution result.\n\nCOPY-AND-PASTE REQUEST\n\n{prompt}\n\nBEGIN_BASE64_NEXT_PHASE_PACKET\n'''
text=art/'Russian_Companion_Complete_Next_Phase_Packet.txt';text.write_text(header+'\n'.join(textwrap.wrap(base64.b64encode(data).decode('ascii'),120))+'\n')
assert base64.b64decode(text.read_text().split('\nBEGIN_BASE64_NEXT_PHASE_PACKET\n',1)[1])==data
shutil.copyfile(n/'SAME_CHAT_PROMPT.md',art/'Russian_Companion_Same_Chat_Next_Prompt.txt')
record={'target':manifest['auditTarget'],'inputAuditArchiveSha256':'199b3229665b769b043434f1d6a69abb242d91b16fd6ed26187fad1acb294185','archive':{'file':archive.name,'bytes':len(data),'sha256':digest},'wordInstructions':{'file':word_path.name,'bytes':word_path.stat().st_size,'sha256':hashlib.sha256(word_path.read_bytes()).hexdigest()},'completeTextPacket':{'file':text.name,'bytes':text.stat().st_size,'sha256':hashlib.sha256(text.read_bytes()).hexdigest(),'losslessArchiveRecoveryVerified':True},'archiveCrcAndAllManifestHashesVerified':True,'actualAuditOrFixesPerformed':False,'runtimeTestsExecuted':0,'privatePdfIncluded':False,'physicalAcceptance':'OUTSTANDING_LAST'}
(n/'DELIVERY.json').write_text(json.dumps(record,indent=2)+'\n');print(json.dumps(record))
